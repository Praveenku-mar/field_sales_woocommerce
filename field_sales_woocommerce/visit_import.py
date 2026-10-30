# Copyright (c) 2026, Your Company and contributors
# For license information, please see license.txt

"""
Bulk-create Visits from an Excel / CSV file.

Flow: the user downloads a template, fills it in and uploads it (a private File);
``preview_visit_import`` checks every row without writing anything; and
``create_visits_from_import`` validates the file again on the server and inserts
the valid rows. Rows are always judged against the calling user's permissions:
a rep may only plan visits for reps in their own hierarchy.
"""

import datetime
from io import BytesIO

import frappe
import xlsxwriter
from frappe import _
from frappe.desk.utils import provide_binary_file
from frappe.utils import getdate, nowdate
from frappe.utils.csvutils import read_csv_content
from frappe.utils.xlsxutils import make_xlsx
from openpyxl import load_workbook

from field_sales_woocommerce.permissions import get_allowed_reps, is_admin

MAX_ROWS = 500
DEFAULT_PURPOSE = "Order Taking"
DATE_FORMATS = ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d")

# Template column -> accepted header spellings (compared lower-case, spaces and underscores alike).
COLUMNS = {
	"sales_rep": ("sales rep", "salesrep", "rep"),
	"customer": ("customer", "customer id", "customer name", "retailer", "retailer id", "retailer name"),
	"visit_date": ("visit date", "date"),
	"purpose": ("purpose", "visit purpose"),
	"notes": ("notes", "note", "remarks"),
}
REQUIRED_COLUMNS = ("sales_rep", "customer", "visit_date")
# Bulk planning is a manager task; reps create visits one at a time.
IMPORT_ROLES = {"System Manager", "Sales Manager"}


def _check_import_access() -> None:
	if frappe.session.user != "Administrator" and not IMPORT_ROLES & set(frappe.get_roles()):
		frappe.throw(_("Only a Sales Manager or System Manager can import visits."), frappe.PermissionError)


# ----------------------------------------------------------------------
# Template
# ----------------------------------------------------------------------


def _get_purposes() -> list[str]:
	return [p for p in frappe.get_meta("Visit").get_field("visit_purpose").options.split("\n") if p]


def _get_reps(active_only: bool = True) -> list[dict]:
	"""Sales Reps the user may plan visits for: any for admins, otherwise their hierarchy."""
	filters = {"active": 1} if active_only else {}
	if not is_admin():
		allowed = get_allowed_reps()
		if not allowed:
			return []
		filters["name"] = ("in", allowed)
	return frappe.get_all(
		"Sales Rep", filters=filters, fields=["name", "full_name", "active"], order_by="full_name asc", limit_page_length=0
	)


def _get_customers(active_only: bool = True) -> list[dict]:
	"""Retailers the user may read (get_list applies Retailer permissions)."""
	filters = {"status": "Active"} if active_only else {}
	return frappe.get_list(
		"Retailer",
		filters=filters,
		fields=["name", "retailer_name", "city", "status"],
		order_by="retailer_name asc",
		limit_page_length=0,
	)


@frappe.whitelist()
def download_visit_template() -> None:
	"""Excel template: a "Visits" sheet to fill in and a "Lists" sheet with the valid reps, customers and purposes."""
	_check_import_access()
	reps = _get_reps()
	customers = _get_customers()
	purposes = _get_purposes()

	example = [
		reps[0]["name"] if reps else "",
		customers[0]["name"] if customers else "",
		nowdate(),
		DEFAULT_PURPOSE,
		"Example row: replace or delete it",
	]
	visits = [["Sales Rep", "Customer", "Visit Date", "Purpose", "Notes"], example]

	lists = [["Sales Rep email", "Sales Rep name", "Customer ID", "Customer name", "City", "Purposes"]]
	for i in range(max(len(reps), len(customers), len(purposes))):
		rep = reps[i] if i < len(reps) else {}
		customer = customers[i] if i < len(customers) else {}
		lists.append(
			[
				rep.get("name", ""),
				rep.get("full_name", ""),
				customer.get("name", ""),
				customer.get("retailer_name", ""),
				customer.get("city", "") or "",
				purposes[i] if i < len(purposes) else "",
			]
		)

	output = BytesIO()
	wb = xlsxwriter.Workbook(output, {"constant_memory": True})
	make_xlsx(visits, "Visits", wb=wb, column_widths=[28, 28, 14, 16, 40])
	make_xlsx(lists, "Lists", wb=wb, column_widths=[28, 24, 16, 32, 18, 16])
	wb.close()
	# The template lists customers and reps (personal data): record who downloaded it.
	from frappe.core.doctype.access_log.access_log import make_access_log

	make_access_log(
		doctype="Retailer",
		method="field_sales_woocommerce.visit_import.download_visit_template",
		file_type="xlsx",
		columns="Sales Rep email, Sales Rep name, Customer ID, Customer name, City",
	)
	provide_binary_file("visit_import_template", "xlsx", output.getvalue())


# ----------------------------------------------------------------------
# Reading the upload
# ----------------------------------------------------------------------


def _get_file(file_url: str):
	"""The caller's own private upload; anything else is refused."""
	if not file_url or not file_url.startswith("/private/files/"):
		frappe.throw(_("Upload the file again."), title=_("Invalid File"))
	name = frappe.db.get_value("File", {"file_url": file_url, "is_private": 1, "owner": frappe.session.user})
	if not name:
		frappe.throw(_("File not found. Upload it again."), title=_("Invalid File"))
	return frappe.get_doc("File", name)


def _read_rows(file) -> list[list]:
	"""All rows of the first sheet (or of the CSV) as lists of raw cell values."""
	ext = (file.file_name or "").rsplit(".", 1)[-1].lower()
	try:
		if ext == "xlsx":
			wb = load_workbook(filename=file.get_full_path(), data_only=True, read_only=True)
			try:
				return [list(row) for row in wb.worksheets[0].iter_rows(values_only=True)]
			finally:
				wb.close()
		if ext == "csv":
			return read_csv_content(file.get_content(), use_sniffer=True)
	except Exception:
		frappe.log_error("Visit import: unreadable file")
		frappe.throw(_("Could not read the file. Use the template (.xlsx) or a CSV."), title=_("Invalid File"))
	frappe.throw(_("Only .xlsx and .csv files can be imported."), title=_("Invalid File"))


def _text(value) -> str:
	if value is None:
		return ""
	if isinstance(value, float) and value.is_integer():
		value = int(value)
	return str(value).strip()


def _map_columns(header: list) -> dict[str, int]:
	"""Column index per field, matching headers case-insensitively."""
	wanted = {}
	for field, spellings in COLUMNS.items():
		wanted.update({s: field for s in spellings})

	mapping = {}
	for idx, cell in enumerate(header):
		key = " ".join(_text(cell).lower().replace("_", " ").split())
		field = wanted.get(key)
		if field and field not in mapping:
			mapping[field] = idx

	missing = [f for f in REQUIRED_COLUMNS if f not in mapping]
	if missing:
		labels = {"sales_rep": "Sales Rep", "customer": "Customer", "visit_date": "Visit Date"}
		frappe.throw(
			_("Missing column(s): {0}. Download the template and use its headings.").format(
				", ".join(labels[f] for f in missing)
			),
			title=_("Invalid File"),
		)
	return mapping


def _parse_date(value) -> str | None:
	"""ISO date from a date cell or DD-MM-YYYY / DD/MM/YYYY / YYYY-MM-DD text; None when unreadable."""
	if isinstance(value, datetime.datetime):
		return value.date().isoformat()
	if isinstance(value, datetime.date):
		return value.isoformat()
	text = _text(value)
	for fmt in DATE_FORMATS:
		try:
			return datetime.datetime.strptime(text, fmt).date().isoformat()
		except ValueError:
			continue
	return None


# ----------------------------------------------------------------------
# Validation
# ----------------------------------------------------------------------


def _validate_file(file_url: str) -> list[dict]:
	"""Resolve and check every row of the upload; each result has ``error`` set (or None when ready)."""
	file = _get_file(file_url)
	rows = _read_rows(file)
	if not rows:
		frappe.throw(_("The file is empty."), title=_("Invalid File"))

	columns = _map_columns(rows[0])
	data = []  # (sheet row number, cells)
	for number, cells in enumerate(rows[1:], start=2):
		if any(_text(c) for c in cells):
			data.append((number, cells))
	if len(data) > MAX_ROWS:
		frappe.throw(
			_("The file has {0} rows. Import at most {1} at a time.").format(len(data), MAX_ROWS),
			title=_("Too Many Rows"),
		)
	if not frappe.has_permission("Visit", "create"):
		frappe.throw(_("You are not allowed to create Visits."), title=_("Not Permitted"), exc=frappe.PermissionError)

	reps_by_name, reps_by_full = {}, {}
	for rep in _get_reps(active_only=False):
		reps_by_name[rep.name.lower()] = rep
		reps_by_full.setdefault((rep.full_name or "").strip().lower(), []).append(rep)
	customers_by_id, customers_by_name = {}, {}
	for customer in _get_customers(active_only=False):
		customers_by_id[customer.name.lower()] = customer
		customers_by_name.setdefault((customer.retailer_name or "").strip().lower(), []).append(customer)

	valid_purposes = {p.lower(): p for p in _get_purposes()}
	seen = {}
	results = []
	for number, cells in data:

		def cell(field):
			idx = columns.get(field)
			return cells[idx] if idx is not None and idx < len(cells) else None

		result = {
			"row": number,
			"sales_rep": None,
			"sales_rep_name": _text(cell("sales_rep")),
			"retailer": None,
			"retailer_name": _text(cell("customer")),
			"city": None,
			"visit_date": None,
			"visit_purpose": None,
			"notes": _text(cell("notes")),
			"error": None,
		}
		result["error"] = _check_row(
			result,
			cell("visit_date"),
			_text(cell("purpose")),
			(reps_by_name, reps_by_full),
			(customers_by_id, customers_by_name),
			valid_purposes,
			seen,
		)
		results.append(result)
	return results


def _check_row(result, date_cell, purpose, reps, customers, valid_purposes, seen) -> str | None:
	"""Fill ``result`` with the resolved values and return the first problem found, if any."""
	reps_by_name, reps_by_full = reps
	customers_by_id, customers_by_name = customers

	rep_text = result["sales_rep_name"]
	if not rep_text:
		return _("Sales Rep is missing")
	rep = reps_by_name.get(rep_text.lower())
	if not rep:
		matches = reps_by_full.get(rep_text.lower(), [])
		if len(matches) > 1:
			return _("Several Sales Reps are named '{0}'. Use the email instead.").format(rep_text)
		rep = matches[0] if matches else None
	if not rep:
		return _("Sales Rep '{0}' not found in your team").format(rep_text)
	result["sales_rep"], result["sales_rep_name"] = rep.name, rep.full_name or rep.name

	customer_text = result["retailer_name"]
	if not customer_text:
		return _("Customer is missing")
	customer = customers_by_id.get(customer_text.lower())
	if not customer:
		matches = customers_by_name.get(customer_text.lower(), [])
		if len(matches) > 1:
			options = ", ".join(f"{c.city or _('no city')}: {c.name}" for c in matches)
			return _("Several customers are named '{0}' ({1}). Use the Customer ID.").format(customer_text, options)
		customer = matches[0] if matches else None
	if not customer:
		return _("Customer '{0}' not found").format(customer_text)
	result["retailer"], result["retailer_name"], result["city"] = customer.name, customer.retailer_name, customer.city

	visit_date = _parse_date(date_cell)
	if not visit_date:
		return _("Visit Date '{0}' is not a date. Use DD-MM-YYYY.").format(_text(date_cell))
	result["visit_date"] = visit_date

	if purpose and purpose.lower() not in valid_purposes:
		return _("Purpose '{0}' is not one of: {1}").format(purpose, ", ".join(valid_purposes.values()))
	result["visit_purpose"] = valid_purposes[purpose.lower()] if purpose else DEFAULT_PURPOSE

	if not rep.active:
		return _("Sales Rep {0} is inactive").format(result["sales_rep_name"])
	if customer.status != "Active":
		return _("Customer {0} is inactive").format(customer.retailer_name)
	if not frappe.has_permission("Retailer", "read", customer.name):
		return _("You cannot access customer {0}").format(customer.name)

	key = (rep.name, customer.name, visit_date)
	if key in seen:
		return _("Duplicate of row {0} in this file").format(seen[key])
	existing = frappe.db.get_value(
		"Visit",
		{"sales_rep": rep.name, "retailer": customer.name, "visit_date": visit_date, "status": ("!=", "Cancelled")},
	)
	if existing:
		return _("Visit {0} already exists for this rep, customer and date").format(existing)
	seen[key] = result["row"]
	return None


# ----------------------------------------------------------------------
# Endpoints
# ----------------------------------------------------------------------


@frappe.whitelist()
def preview_visit_import(file_url: str) -> dict:
	"""Check an uploaded file row by row without creating anything."""
	_check_import_access()
	rows = _validate_file(file_url)
	errors = sum(1 for r in rows if r["error"])
	return {"rows": rows, "ready": len(rows) - errors, "errors": errors}


@frappe.whitelist(methods=["POST"])
def create_visits_from_import(file_url: str) -> dict:
	"""
	Validate the file again (the preview is never trusted) and insert a Planned Visit
	for each valid row. Each row has its own savepoint, so one failure leaves the rest alone.

	The upload is a throwaway working copy: once at least one visit is created it is
	deleted, so private spreadsheets of customers and reps do not pile up. When nothing
	was created it is kept, so the user can fix the data and try again.
	"""
	_check_import_access()
	rows = _validate_file(file_url)
	created, skipped = [], []
	for row in rows:
		if row["error"]:
			skipped.append({"row": row["row"], "error": row["error"]})
			continue
		frappe.db.savepoint("visit_import_row")
		try:
			visit = frappe.get_doc(
				{
					"doctype": "Visit",
					"sales_rep": row["sales_rep"],
					"retailer": row["retailer"],
					"visit_date": row["visit_date"],
					"visit_purpose": row["visit_purpose"],
					"notes": row["notes"],
					"status": "Planned",
				}
			).insert()
		except Exception as e:
			frappe.db.rollback(save_point="visit_import_row")
			skipped.append({"row": row["row"], "error": frappe.utils.strip_html(str(e)) or _("Could not create the visit")})
			continue
		created.append(
			{
				"row": row["row"],
				"name": visit.name,
				"retailer_name": visit.retailer_name,
				"visit_date": str(getdate(visit.visit_date)),
			}
		)

	if created:
		file = _get_file(file_url)
		frappe.delete_doc("File", file.name, ignore_permissions=True, force=True)
	return {"created": created, "skipped": skipped}
